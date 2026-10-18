"""
Genera los documentos binarios de ejemplo de TechRetail (Word, Excel, PowerPoint
y PDF). Los de texto (Markdown, HTML, CSV y JSON) están en el repositorio como
archivos comunes, para poder leerlos y editarlos directamente.

Todo el contenido es ficticio pero coherente entre sí: misma tienda, mismos
productos, mismas políticas.

Uso:
    python -m scripts.generar_documentos
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.config import BASE_DIR  # noqa: E402

# Siempre en el perfil de TechRetail, aunque PERFIL apunte a otra empresa.
DOCS = BASE_DIR / "perfiles" / "techretail" / "documentos"

PRODUCTOS = [
    # (sku, producto, precio de lista en ARS)
    ("NB-156-R5", 'Notebook 15,6" Ryzen 5 16 GB 512 GB SSD', 1_150_000),
    ("NB-140-I5", 'Notebook 14" Core i5 8 GB 256 GB SSD', 890_000),
    ("CEL-65-128", 'Celular 6,5" 128 GB', 420_000),
    ("MON-27-165", 'Monitor 27" IPS 165 Hz', 480_000),
    ("AUR-ANC", "Auriculares inalámbricos con cancelación de ruido", 145_000),
    ("TEC-MEC", "Teclado mecánico inalámbrico", 95_000),
    ("MOU-GAM", "Mouse gamer inalámbrico", 68_000),
    ("CAR-65W", "Cargador USB-C 65 W", 42_000),
]


def _dir(categoria: str) -> Path:
    d = DOCS / categoria
    d.mkdir(parents=True, exist_ok=True)
    return d


# ==========================================================================
#  DOCX (Word)
# ==========================================================================
def gen_docx():
    from docx import Document

    # Política de vacaciones (RR. HH.)
    doc = Document()
    doc.add_heading("Política de vacaciones y licencias - TechRetail", level=1)
    doc.add_paragraph(
        "Uso interno. Regula las vacaciones y licencias de todos los colaboradores "
        "de TechRetail S.R.L."
    )
    doc.add_heading("Vacaciones anuales", level=2)
    doc.add_paragraph(
        "Los días de vacaciones se otorgan según la antigüedad, conforme a la Ley "
        "de Contrato de Trabajo (LCT):"
    )
    tabla = doc.add_table(rows=1, cols=2)
    tabla.style = "Light Grid Accent 1"
    hdr = tabla.rows[0].cells
    hdr[0].text = "Antigüedad"
    hdr[1].text = "Días corridos"
    for antig, dias in [
        ("Menos de 5 años", "14 días"),
        ("De 5 a 10 años", "21 días"),
        ("De 10 a 20 años", "28 días"),
        ("Más de 20 años", "35 días"),
    ]:
        fila = tabla.add_row().cells
        fila[0].text = antig
        fila[1].text = dias
    doc.add_paragraph(
        "Además, la empresa otorga 3 días hábiles extra por año a partir del segundo "
        "año de antigüedad."
    )
    doc.add_heading("Licencias especiales", level=2)
    doc.add_paragraph("• Licencia por matrimonio: 10 días corridos.")
    doc.add_paragraph("• Licencia por paternidad: 15 días corridos.")
    doc.add_paragraph("• Licencia por examen: hasta 4 días por año.")
    doc.add_heading("Solicitud", level=2)
    doc.add_paragraph(
        "Las vacaciones se piden con al menos 30 días de anticipación en el sistema "
        "de gestión y las aprueba el encargado del área. En el local y el depósito no "
        "se otorgan vacaciones del 15 de noviembre al 6 de enero (temporada alta)."
    )
    doc.save(str(_dir("rrhh") / "politica_vacaciones_y_licencias.docx"))

    # Procedimiento de conciliación (Operaciones)
    doc2 = Document()
    doc2.add_heading("Procedimiento de conciliación de cobros - TechRetail", level=1)
    doc2.add_paragraph(
        "Uso interno. Describe cómo se controla todos los días que lo que liquidan "
        "Mercado Pago y Mercado Libre coincida con las ventas registradas."
    )
    doc2.add_heading("Pasos", level=2)
    doc2.add_paragraph(
        "1. Cada mañana, Administración descarga las liquidaciones del día anterior "
        "de Mercado Pago y de Mercado Libre."
    )
    doc2.add_paragraph(
        "2. Se cruzan contra las ventas del sistema de gestión, descontando las "
        "comisiones de la tabla de comisiones vigente."
    )
    doc2.add_paragraph(
        "3. Toda venta con una diferencia mayor al 0,5 % entre lo esperado y lo "
        "liquidado se marca como discrepancia y se revisa en el día."
    )
    doc2.add_paragraph(
        "4. Las devoluciones y los contracargos se registran aparte y se informan al "
        "encargado de postventa."
    )
    doc2.save(str(_dir("operacional") / "procedimiento_conciliacion.docx"))


# ==========================================================================
#  XLSX (Excel)
# ==========================================================================
def gen_xlsx():
    from openpyxl import Workbook

    # Lista de precios (Comercial)
    wb = Workbook()
    ws = wb.active
    ws.title = "Precios"
    ws.append(["sku", "producto", "precio_lista_ars", "precio_transferencia_ars", "uso"])
    for sku, producto, precio in PRODUCTOS:
        ws.append([sku, producto, precio, round(precio * 0.9), "interno"])
    wb.save(str(_dir("comercial") / "lista_de_precios.xlsx"))

    # Ventas por canal (Finanzas)
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "Ventas"
    ws2.append(["mes", "canal", "pedidos", "facturacion_ars"])
    datos = [
        ["Julio 2026", "Tienda online", 412, 186_300_000],
        ["Julio 2026", "Mercado Libre", 938, 301_500_000],
        ["Julio 2026", "Local", 265, 97_800_000],
        ["Agosto 2026", "Tienda online", 455, 204_100_000],
        ["Agosto 2026", "Mercado Libre", 1_012, 327_900_000],
        ["Agosto 2026", "Local", 248, 92_400_000],
        ["Septiembre 2026", "Tienda online", 498, 226_700_000],
        ["Septiembre 2026", "Mercado Libre", 1_096, 352_600_000],
        ["Septiembre 2026", "Local", 271, 101_200_000],
    ]
    for fila in datos:
        ws2.append(fila)
    wb2.save(str(_dir("financiero") / "ventas_por_canal.xlsx"))


# ==========================================================================
#  PPTX (PowerPoint)
# ==========================================================================
def gen_pptx():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "TechRetail S.R.L."
    slide.placeholders[1].text = "Ventas corporativas: equipamiento tecnológico para empresas"

    def slide_bullets(titulo, bullets, notas=""):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = titulo
        cuerpo = s.placeholders[1].text_frame
        cuerpo.text = bullets[0]
        for b in bullets[1:]:
            cuerpo.add_paragraph().text = b
        if notas:
            s.notes_slide.notes_text_frame.text = notas

    slide_bullets(
        "Qué ofrecemos",
        [
            "Notebooks, monitores, celulares y accesorios para equipos de trabajo.",
            "Entrega con configuración inicial: sistema operativo, cuentas y programas.",
            "Factura A y cuenta corriente para clientes con más de 6 meses de compras.",
        ],
    )
    slide_bullets(
        "Descuentos por volumen",
        [
            "5 % desde 10 unidades del mismo producto.",
            "8 % desde 30 unidades.",
            "Los descuentos no se suman al de pago por transferencia.",
        ],
        notas="Pedidos de más de 50 unidades: cotizar con el encargado comercial.",
    )
    slide_bullets(
        "Garantía y soporte",
        [
            "Garantía oficial del fabricante de 12 meses.",
            "Garantía extendida opcional de 12 meses más: 10 % del precio del equipo.",
            "Equipo de reemplazo mientras dura la reparación, en contratos de más de 20 equipos.",
        ],
    )
    prs.save(str(_dir("comercial") / "presentacion_ventas_corporativas.pptx"))


# ==========================================================================
#  PDF
# ==========================================================================
def gen_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    estilos = getSampleStyleSheet()

    def construir(ruta, titulo, parrafos):
        doc = SimpleDocTemplate(str(ruta), pagesize=A4)
        flujo = [Paragraph(titulo, estilos["Title"]), Spacer(1, 12)]
        for p in parrafos:
            if p.startswith("## "):
                flujo.append(Paragraph(p[3:], estilos["Heading2"]))
            else:
                flujo.append(Paragraph(p, estilos["BodyText"]))
            flujo.append(Spacer(1, 6))
        doc.build(flujo)

    # Términos y condiciones (público)
    construir(
        _dir("clientes") / "terminos_y_condiciones.pdf",
        "Términos y condiciones de compra - TechRetail",
        [
            "## 1. Precios",
            "Los precios están expresados en pesos argentinos e incluyen IVA. Pueden "
            "cambiar sin aviso; se respeta el precio vigente al momento de la compra.",
            "## 2. Stock",
            "Las ventas están sujetas a disponibilidad. Si un producto se agota después "
            "de la compra, se ofrece otro equivalente o el reintegro total.",
            "## 3. Imágenes",
            "Las fotos son ilustrativas. Las características que valen son las de la "
            "ficha técnica del producto.",
            "## 4. Cambios, devoluciones y garantía",
            "Se rigen por las políticas de cambios y devoluciones y de garantía "
            "publicadas en la tienda online, y por la Ley de Defensa del Consumidor.",
            "## 5. Jurisdicción",
            "Ante cualquier conflicto, son competentes los tribunales ordinarios de la "
            "ciudad de Rosario, sin perjuicio de los derechos del consumidor.",
        ],
    )

    # Manual operativo (Operaciones)
    construir(
        _dir("operacional") / "manual_procesos_operativos.pdf",
        "Manual de procesos operativos - TechRetail",
        [
            "Uso interno.",
            "## Recepción de mercadería",
            "Se controla el remito contra la orden de compra, se registran los números "
            "de serie de cada equipo y se carga el ingreso en el sistema de gestión antes "
            "de dejarlo disponible para la venta.",
            "## Preparación de pedidos",
            "Al preparar un pedido se controla que el número de serie coincida con el "
            "registrado en la venta, que el equipo tenga todos sus accesorios y que la "
            "caja esté sellada. Los equipos viajan con doble caja y film de burbujas.",
            "## Despacho",
            "Los envíos en el día de Rosario salen a las 14 h. Los envíos por correo se "
            "entregan al transporte a las 16 h; lo que se prepara después sale al día "
            "hábil siguiente.",
            "## Garantías",
            "Cada equipo en garantía se registra en el sistema con el número de pedido y "
            "la falla informada. El servicio técnico lo prueba dentro de las 48 horas; si "
            "corresponde, se envía al servicio oficial del fabricante y se hace el "
            "seguimiento una vez por semana.",
            "## Inventario",
            "Los productos de alta rotación (celulares, cargadores y auriculares) se "
            "cuentan todas las semanas; el resto, una vez por mes.",
        ],
    )


def main():
    DOCS.mkdir(parents=True, exist_ok=True)
    gen_docx()
    gen_xlsx()
    gen_pptx()
    gen_pdf()

    archivos = sorted(p for p in DOCS.rglob("*") if p.is_file())
    print(f"{len(archivos)} documentos en {DOCS}:")
    for a in archivos:
        print(f"  - {a.relative_to(DOCS).as_posix()}")


if __name__ == "__main__":
    main()
